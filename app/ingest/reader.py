"""Turn an uploaded file (CSV or Excel) into headers + rows of strings.

Every value comes out as a string ("" for blanks) so the rest of the
pipeline handles one simple shape no matter what the client sent.
"""

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl import load_workbook


class UnsupportedFileError(ValueError):
    pass


@dataclass
class ParsedFile:
    headers: list[str]
    rows: list[dict[str, str]]  # each row: {original header -> raw string value}


def read_file(filename: str, content: bytes) -> ParsedFile:
    name = filename.lower()
    if name.endswith(".csv"):
        return _read_csv(content)
    if name.endswith(".xlsx"):
        return _read_xlsx(content)
    raise UnsupportedFileError("Only .csv and .xlsx files are supported")


def _read_csv(content: bytes) -> ParsedFile:
    # utf-8-sig strips the invisible BOM that Excel adds when saving CSVs.
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    all_rows = [r for r in reader]
    if not all_rows:
        return ParsedFile(headers=[], rows=[])
    headers = all_rows[0]
    rows = [_zip_row(headers, r) for r in all_rows[1:]]
    return ParsedFile(headers=headers, rows=rows)


def _read_xlsx(content: bytes) -> ParsedFile:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    values = [[_cell_to_str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    if not values:
        return ParsedFile(headers=[], rows=[])
    headers = values[0]
    rows = [_zip_row(headers, r) for r in values[1:]]
    return ParsedFile(headers=headers, rows=rows)


def _zip_row(headers: list[str], values: list[str]) -> dict[str, str]:
    # Short rows get "" for the missing cells; extra cells are dropped.
    padded = list(values) + [""] * (len(headers) - len(values))
    return {h: padded[i] for i, h in enumerate(headers)}


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # 9876543210.0 -> "9876543210"
    return str(value)
